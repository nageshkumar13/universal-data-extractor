from copy import deepcopy
import json

from pathlib import Path
from unittest.mock import Mock

import pytest
import yaml

from openpyxl import load_workbook

import core.runner as runner_module
from core.runner import ScrapeRunner


class FakeLoader:
    def __init__(self, profile: dict):
        self.profile = profile

    def load(self, path: Path) -> dict:
        return self.profile


class FakeClient:
    def __init__(self, html_by_url: dict[str, str]):
        self.html_by_url = html_by_url
        self.calls: list[str] = []
        self.enter_count = 0
        self.exit_count = 0
        self.closed = False
        self.exit_exception: BaseException | None = None

    def __enter__(self):
        self.enter_count += 1
        self.closed = False
        return self

    def __exit__(self, exc_type, exc_value, traceback):
        self.exit_count += 1
        self.exit_exception = exc_value
        self.close()

    def close(self):
        self.closed = True

    def fetch(self, url: str, wait_for: str | None = None) -> str:
        assert self.enter_count == 1 and self.exit_count == 0
        self.calls.append(url)
        return self.html_by_url[url]


class FakeParser:
    def __init__(self, records_by_url: dict[str, list[dict]]):
        self.records_by_url = records_by_url

    def extract(self, html: str, fields: dict[str, str], base_url: str) -> list[dict]:
        return self.records_by_url[base_url]


class FakePaginator:
    def __init__(self, next_urls: dict[str, str | None]):
        self.next_urls = next_urls

    def get_next_url(self, html: str, current_url: str, next_selector: str) -> str | None:
        return self.next_urls[current_url]


class FakeCache:
    def __init__(self, cached_urls: set[str] | None = None):
        self.cached_urls = set(cached_urls or set())
        self.marked: list[tuple[str, int]] = []
        self.cleared = False
        self.completions = {}
        self.completed = []

    def is_cached(self, url: str) -> bool:
        return url in self.cached_urls

    def mark_done(self, url: str, record_count: int) -> None:
        self.cached_urls.add(url)
        self.marked.append((url, record_count))

    def clear(self) -> None:
        self.cached_urls.clear()
        self.completions.clear()
        self.cleared = True


    def completed_pages(self, output_key, fingerprint):
        entry = self.completions.get(output_key)
        return entry[1] if entry is not None and entry[0] == fingerprint else None

    def invalidate_run(self, output_key):
        self.completions.pop(output_key, None)

    def mark_complete(self, output_key, fingerprint, pages_scraped):
        self.completions[output_key] = (fingerprint, pages_scraped)
        self.completed.append((output_key, fingerprint, pages_scraped))

    def clear_runs(self):
        self.completions.clear()
        self.cleared = True


class FakeTransformer:
    def transform(self, records: list[dict]) -> list[dict]:
        return records


class FakeExporter:
    def write_quality_artifacts(self, result, output_path):
        return runner_module.Exporter().write_quality_artifacts(result, output_path)

    def __init__(self) -> None:
        self.calls: list[tuple[str, list[dict], Path]] = []

    def to_csv(self, records: list[dict], output_path: Path) -> Path:
        self.calls.append(("csv", list(records), output_path))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text("csv", encoding="utf-8")
        return output_path

    def to_json(self, records: list[dict], output_path: Path) -> Path:
        self.calls.append(("json", list(records), output_path))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text("json", encoding="utf-8")
        return output_path

    def to_excel(self, records: list[dict], output_path: Path) -> Path:
        self.calls.append(("xlsx", list(records), output_path))
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text("xlsx", encoding="utf-8")
        return output_path


class AllowAllRobots:
    def __init__(self, base_url: str):
        self.base_url = base_url

    def is_allowed(self, url: str, user_agent: str = "*") -> bool:
        return True


def test_completed_run_preserves_outputs_without_opening_client(tmp_path, make_quality_runner):
    runner, path = make_quality_runner(quality=False)
    first = runner.run(path, tmp_path)
    before = {key: first[key].read_bytes() for key in ("csv_path", "json_path", "xlsx_path")}
    runner.client = FakeClient({})
    runner.exporter.calls.clear()

    summary = runner.run(path, tmp_path)

    assert summary["cache_only_run"] is True
    assert summary["pages_scraped"] == 0
    assert summary["cached_skips"] == 1
    assert runner.client.enter_count == runner.client.exit_count == 0
    assert runner.client.calls == []
    assert runner.exporter.calls == []
    assert {key: summary[key].read_bytes() for key in before} == before


def test_runner_respects_profile_delay_between_requests(tmp_path, monkeypatch) -> None:
    page_one = "https://example.com/page-1.html"
    page_two = "https://example.com/page-2.html"
    profile = {
        "site_name": "Delayed Example",
        "engine": "static",
        "start_url": page_one,
        "fields": {"title": "article h2::text"},
        "max_pages": 2,
        "delay": 1.0,
        "pagination": {"next_button": "a.next"},
    }
    sleeps: list[float] = []

    runner = ScrapeRunner()
    runner.loader = FakeLoader(profile)
    runner.client = FakeClient(
        {
            page_one: "<html>page one</html>",
            page_two: "<html>page two</html>",
        }
    )
    runner.parser = FakeParser(
        {
            page_one: [{"title": "One"}],
            page_two: [{"title": "Two"}],
        }
    )
    runner.paginator = FakePaginator(
        {
            page_one: page_two,
            page_two: None,
        }
    )
    runner.cache = FakeCache()
    runner.transformer = FakeTransformer()
    runner.exporter = FakeExporter()

    monkeypatch.setattr(runner_module, "RobotsChecker", AllowAllRobots)
    monkeypatch.setattr(runner_module.time, "sleep", lambda seconds: sleeps.append(seconds))

    stages = []
    original_transform = runner.transformer.transform

    def transform_after_exit(records):
        assert runner.client.exit_count == 1
        assert runner.client.closed is True
        stages.append("transform")
        return original_transform(records)

    monkeypatch.setattr(runner.transformer, "transform", transform_after_exit)

    def check_export_after_exit(method, stage):
        def export(records, output_path):
            assert runner.client.exit_count == 1
            assert runner.client.closed is True
            stages.append(stage)
            return method(records, output_path)
        return export

    for method_name in ("to_csv", "to_json", "to_excel"):
        method = getattr(runner.exporter, method_name)
        monkeypatch.setattr(
            runner.exporter, method_name, check_export_after_exit(method, method_name),
        )

    summary = runner.run(Path("profiles/example.yaml"), tmp_path)

    assert summary["pages_scraped"] == 2
    assert runner.client.calls == [page_one, page_two]
    assert sleeps == [1.0]
    assert runner.client.enter_count == runner.client.exit_count == 1
    assert runner.client.exit_exception is None
    assert stages == ["transform", "to_csv", "to_json", "to_excel"]
    assert summary["records_extracted"] == summary["records_transformed"] == 2
    assert summary["cached_skips"] == summary["robots_blocked"] == 0
    assert runner.cache.marked == []
    assert len(runner.cache.completed) == 1
    assert runner.cache.completed[0][2] == 2


@pytest.mark.parametrize("error_type", [RuntimeError, KeyboardInterrupt])
@pytest.mark.parametrize("client_source", ["injected", "factory"])
def test_runner_exits_client_when_fetch_raises(
    tmp_path, monkeypatch, error_type, client_source,
):
    profile = {
        "site_name": "Failing Example",
        "engine": "browser",
        "start_url": "https://example.com/page-1.html",
        "fields": {"title": "article h2::text"},
        "max_pages": 1,
        "delay": 0,
    }
    client = FakeClient({})
    error = error_type("fetch failed")

    def failing_fetch(url, wait_for=None):
        assert client.enter_count == 1 and client.exit_count == 0
        raise error

    monkeypatch.setattr(client, "fetch", failing_fetch)
    factory = Mock(return_value=client)
    monkeypatch.setattr(runner_module, "create_client", factory)
    monkeypatch.setattr(runner_module, "RobotsChecker", AllowAllRobots)
    runner = ScrapeRunner(client=client if client_source == "injected" else None)
    runner.loader = FakeLoader(profile)
    runner.cache = FakeCache()
    runner.transformer = Mock(spec=["transform"])
    runner.exporter = FakeExporter()

    with pytest.raises(error_type) as caught:
        runner.run(Path("profiles/example.yaml"), tmp_path)

    assert caught.value is error
    assert client.enter_count == client.exit_count == 1
    assert client.exit_exception is error
    assert client.closed is True
    runner.transformer.transform.assert_not_called()
    assert runner.exporter.calls == []
    if client_source == "factory":
        factory.assert_called_once_with("browser")
    else:
        factory.assert_not_called()


@pytest.fixture
def make_quality_runner(tmp_path, monkeypatch):
    monkeypatch.setattr(runner_module, "URLCache", FakeCache)
    monkeypatch.setattr(runner_module, "RobotsChecker", AllowAllRobots)

    def make(quality=True, html_by_url=None):
        start_url = "https://example.com/page-1.html"
        selectors = {
            "id": "article::attr(data-id)",
            "title": "article h2::text",
            "price": "article .price::text",
        }
        profile = {
            "site_name": "Quality Example",
            "engine": "static",
            "start_url": start_url,
            "fields": selectors,
            "max_pages": 2,
            "delay": 0,
            "pagination": {"next_button": "a.next"},
        }
        if quality is not None:
            profile["data_quality"] = quality
        if quality is True:
            profile.update({
                "fields": {
                    "id": {"selector": selectors["id"], "required": True, "type": "int"},
                    "title": selectors["title"],
                    "price": {"selector": selectors["price"], "type": "float"},
                },
                "unique_key": ["id"],
            })
        profile_path = tmp_path / "integration.yaml"
        profile_path.write_text(yaml.safe_dump(profile, sort_keys=False), encoding="utf-8")
        if html_by_url is None:
            html_by_url = {
                start_url: '<article data-id="1"><h2>One</h2><p class="price">12.50</p></article>',
            }
        runner = ScrapeRunner(client=FakeClient(html_by_url))
        runner.exporter = FakeExporter()
        return runner, profile_path

    return make


@pytest.mark.parametrize("quality", [None, False])
def test_legacy_run_skips_quality_and_preserves_export_and_return_contract(
    tmp_path, make_quality_runner, quality,
):
    runner, path = make_quality_runner(quality=quality)
    runner.quality_processor = Mock()
    assert runner.last_quality_result is None

    summary = runner.run(path, tmp_path)

    runner.quality_processor.process.assert_not_called()
    assert runner.last_quality_result is None
    expected_records = [{"id": "1", "title": "One", "price": 12.5}]
    assert [call[0] for call in runner.exporter.calls] == ["csv", "json", "xlsx"]
    assert all(call[1] == expected_records for call in runner.exporter.calls)
    assert summary == {
        "site_name": "Quality Example", "pages_scraped": 1,
        "records_extracted": 1, "records_transformed": 1,
        "cached_skips": 0, "robots_blocked": 0, "cache_only_run": False,
        "csv_path": tmp_path / "quality_example.csv",
        "json_path": tmp_path / "quality_example.json",
        "xlsx_path": tmp_path / "quality_example.xlsx",
    }


def test_quality_runs_once_after_transform_across_pages_before_all_exports(
    tmp_path, make_quality_runner,
):
    page_one = "https://example.com/page-1.html"
    page_two = "https://example.com/page-2.html"
    html_by_url = {
        page_one: (
            '<article data-id="01"><h2>First</h2><p class="price">12.50</p></article>'
            '<article data-id="bad"><h2>Rejected</h2><p class="price">1.00</p></article>'
            '<a class="next" href="page-2.html">Next</a>'
        ),
        page_two: (
            '<article data-id="1"><h2>Duplicate</h2><p class="price">15.50</p></article>'
            '<article data-id="2"><h2>Second</h2><p class="price">20.00</p></article>'
        ),
    }
    runner, path = make_quality_runner(html_by_url=html_by_url)
    expected_profile = runner.loader.load(path)
    runner.loader.load = Mock(return_value=expected_profile)
    original_profile = deepcopy(expected_profile)
    stages, parser_batches, original_batches = [], [], []
    original_extract = runner.parser.extract

    def extract(html, fields, base_url):
        assert fields == {
            name: definition["selector"]
            for name, definition in expected_profile["fields"].items()
        }
        records = original_extract(html, fields, base_url)
        parser_batches.append(records)
        original_batches.append(deepcopy(records))
        return records

    runner.parser.extract = Mock(side_effect=extract)
    original_transform = runner.transformer.transform

    def transform(records):
        assert runner.client.closed
        assert records == [record for batch in original_batches for record in batch]
        stages.append("transform")
        return original_transform(records)

    runner.transformer.transform = Mock(side_effect=transform)
    original_process = runner.quality_processor.process
    transformed_snapshots = []

    def process(records, profile):
        assert stages == ["transform"]
        assert runner.client.closed
        assert profile is expected_profile
        assert all(isinstance(record["price"], float) for record in records)
        transformed_snapshots.append(deepcopy(records))
        stages.append("quality")
        return original_process(records, profile)

    runner.quality_processor.process = Mock(side_effect=process)

    def observe_export(method, stage):
        def export(records, output_path):
            assert "quality" in stages
            assert records is runner.last_quality_result.clean_records
            stages.append(stage)
            return method(records, output_path)
        return export

    for name in ("to_csv", "to_json", "to_excel"):
        setattr(runner.exporter, name, observe_export(getattr(runner.exporter, name), name))

    summary = runner.run(path, tmp_path)

    runner.quality_processor.process.assert_called_once()
    result = runner.last_quality_result
    assert result is not None
    assert result.clean_records == [
        {"id": 1, "title": "First", "price": 12.5},
        {"id": 2, "title": "Second", "price": 20.0},
    ]
    assert result.rejected_records[0]["normalized"]["title"] == "Rejected"
    assert result.report.extracted == result.report.valid_before_deduplication + result.report.rejected == 4
    assert result.report.exported == result.report.valid_before_deduplication - result.report.duplicates_removed == 2
    assert stages == ["transform", "quality", "to_csv", "to_json", "to_excel"]
    assert all(call[1] == result.clean_records for call in runner.exporter.calls)
    assert summary["records_extracted"] == summary["records_transformed"] == 4
    assert summary["pages_scraped"] == 2
    assert summary["cached_skips"] == summary["robots_blocked"] == 0
    assert runner.client.calls == [page_one, page_two]
    assert runner.cache.marked == []
    assert len(runner.cache.completed) == 1
    assert runner.cache.completed[0][2] == 2
    assert runner.client.enter_count == runner.client.exit_count == 1
    assert expected_profile == original_profile
    assert parser_batches == original_batches
    assert runner.quality_processor.process.call_args.args[0] == transformed_snapshots[0]
    assert html_by_url == runner.client.html_by_url


@pytest.mark.parametrize("preexisting", [False, True])
def test_all_rejected_records_export_empty_clean_lists(tmp_path, make_quality_runner, preexisting):
    runner, path = make_quality_runner(html_by_url={
        "https://example.com/page-1.html": (
            '<article data-id="bad"><h2>Rejected</h2><p class="price">1.00</p></article>'
        ),
    })
    if preexisting:
        for extension in ("csv", "json", "xlsx"):
            (tmp_path / f"quality_example.{extension}").write_text("previous", encoding="utf-8")
    runner.quality_processor.process = Mock(wraps=runner.quality_processor.process)

    summary = runner.run(path, tmp_path)

    runner.quality_processor.process.assert_called_once()
    assert [call[0] for call in runner.exporter.calls] == ["csv", "json", "xlsx"]
    assert all(call[1] == [] for call in runner.exporter.calls)
    assert runner.last_quality_result.report.rejected == 1
    assert runner.last_quality_result.report.exported == 0
    assert summary["records_transformed"] == 1
    assert summary["cache_only_run"] is False


def test_existing_exporters_accept_empty_quality_output(tmp_path, make_quality_runner):
    runner, path = make_quality_runner(html_by_url={
        "https://example.com/page-1.html": '<article data-id="bad"><h2>Rejected</h2></article>',
    })
    runner.exporter = runner_module.Exporter()

    summary = runner.run(path, tmp_path)

    assert runner.last_quality_result.clean_records == []
    assert summary["csv_path"].read_text(encoding="utf-8").strip() == ""
    assert json.loads(summary["json_path"].read_text(encoding="utf-8")) == []
    workbook = load_workbook(summary["xlsx_path"])
    try:
        assert list(workbook.active.values) == []
    finally:
        workbook.close()


@pytest.mark.parametrize("error_type", [RuntimeError, KeyboardInterrupt])
def test_processor_failure_propagates_without_export_or_stale_result(
    tmp_path, make_quality_runner, error_type,
):
    runner, path = make_quality_runner()
    runner.run(path, tmp_path)
    assert runner.last_quality_result is not None
    runner.client = FakeClient(runner.client.html_by_url)
    runner.exporter.calls.clear()
    error = error_type("quality failed")

    def fail(records, profile):
        assert runner.last_quality_result is None
        assert runner.client.closed
        raise error

    runner.quality_processor.process = Mock(side_effect=fail)

    with pytest.raises(error_type) as caught:
        runner.run(path, tmp_path, clear_cache=True)

    assert caught.value is error
    runner.quality_processor.process.assert_called_once()
    assert runner.exporter.calls == []
    assert runner.last_quality_result is None


@pytest.mark.parametrize("stage", ["load", "fetch", "transform"])
def test_result_is_reset_before_failures_earlier_in_the_pipeline(
    tmp_path, make_quality_runner, stage,
):
    runner, path = make_quality_runner()
    runner.run(path, tmp_path)
    assert runner.last_quality_result is not None
    runner.client = FakeClient(runner.client.html_by_url)
    runner.exporter.calls.clear()
    runner.quality_processor.process = Mock()
    error = RuntimeError(f"{stage} failed")

    def fail(*args, **kwargs):
        assert runner.last_quality_result is None
        raise error

    target, method = {
        "load": (runner.loader, "load"),
        "fetch": (runner.client, "fetch"),
        "transform": (runner.transformer, "transform"),
    }[stage]
    setattr(target, method, Mock(side_effect=fail))

    with pytest.raises(RuntimeError) as caught:
        runner.run(path, tmp_path, clear_cache=True)

    assert caught.value is error
    assert runner.last_quality_result is None
    runner.quality_processor.process.assert_not_called()
    assert runner.exporter.calls == []


@pytest.mark.parametrize("quality", [None, False])
def test_later_legacy_run_clears_previous_quality_result(tmp_path, make_quality_runner, quality):
    runner, path = make_quality_runner()
    runner.run(path, tmp_path)
    assert runner.last_quality_result is not None
    legacy_runner, legacy_path = make_quality_runner(quality=quality)
    runner.client = legacy_runner.client
    runner.quality_processor.process = Mock()
    runner.exporter.calls.clear()
    original_load = runner.loader.load

    def load(profile_path):
        assert runner.last_quality_result is None
        return original_load(profile_path)

    runner.loader.load = Mock(side_effect=load)

    summary = runner.run(legacy_path, tmp_path, clear_cache=True)

    assert runner.last_quality_result is None
    runner.quality_processor.process.assert_not_called()
    assert all(call[1] == [{"id": "1", "title": "One", "price": 12.5}] for call in runner.exporter.calls)
    assert summary["records_transformed"] == 1


def test_completed_quality_rerun_skips_processing_and_resets_result(tmp_path, make_quality_runner):
    runner, path = make_quality_runner()
    runner.run(path, tmp_path)
    assert runner.last_quality_result is not None
    runner.client = FakeClient({})
    runner.exporter.calls.clear()
    runner.quality_processor.process = Mock(wraps=runner.quality_processor.process)

    summary = runner.run(path, tmp_path)

    runner.quality_processor.process.assert_not_called()
    assert runner.last_quality_result is None
    assert runner.exporter.calls == []
    assert runner.client.calls == []
    assert summary["cache_only_run"] is True
    assert summary["cached_skips"] == 1


def test_quality_run_preserves_robots_blocking(tmp_path, make_quality_runner, monkeypatch):
    runner, path = make_quality_runner()
    robots = Mock()
    robots.is_allowed.return_value = False
    monkeypatch.setattr(runner_module, "RobotsChecker", Mock(return_value=robots))
    runner.quality_processor.process = Mock(wraps=runner.quality_processor.process)

    summary = runner.run(path, tmp_path)

    assert summary["robots_blocked"] == 1
    assert summary["pages_scraped"] == summary["records_extracted"] == 0
    assert runner.client.calls == []
    assert runner.cache.marked == []
    runner.quality_processor.process.assert_called_once()
    assert runner.quality_processor.process.call_args.args[0] == []
    assert all(call[1] == [] for call in runner.exporter.calls)


@pytest.mark.parametrize("existing_formats", [(), ("csv",), ("csv", "json", "xlsx")])
def test_zero_extracted_records_produce_zero_quality_result_and_preserve_export_path(
    tmp_path, make_quality_runner, existing_formats,
):
    from dataclasses import asdict

    runner, path = make_quality_runner(html_by_url={
        "https://example.com/page-1.html": "<html><body>No listings</body></html>",
    })
    profile = runner.loader.load(path)
    runner.quality_processor.process = Mock(wraps=runner.quality_processor.process)
    for extension in existing_formats:
        (tmp_path / f"quality_example.{extension}").write_text("previous", encoding="utf-8")

    summary = runner.run(path, tmp_path)

    runner.quality_processor.process.assert_called_once_with([], profile)
    result = runner.last_quality_result
    assert result is not None
    assert result.clean_records == result.rejected_records == []
    assert all(count == 0 for count in asdict(result.report).values())
    result.report.enforce_invariants()
    assert summary["pages_scraped"] == 1
    assert summary["records_extracted"] == summary["records_transformed"] == 0
    assert summary["cache_only_run"] is False
    assert [call[0] for call in runner.exporter.calls] == ["csv", "json", "xlsx"]
    assert all(call[1] == [] for call in runner.exporter.calls)
    assert all(summary[f"{extension}_path"].exists() for extension in ("csv", "json", "xlsx"))



@pytest.mark.parametrize("failing_method", ["to_csv", "to_json", "to_excel"])
def test_exporter_failure_retains_completed_quality_result(
    tmp_path, make_quality_runner, failing_method,
):
    runner, path = make_quality_runner()
    runner.run(path, tmp_path)
    previous_result = runner.last_quality_result
    assert previous_result is not None
    runner.client = FakeClient(runner.client.html_by_url)
    runner.exporter.calls.clear()
    completed_results = []
    original_process = runner.quality_processor.process

    def process(records, profile):
        assert runner.last_quality_result is None
        result = original_process(records, profile)
        completed_results.append(result)
        return result

    runner.quality_processor.process = Mock(side_effect=process)
    error = OSError(f"{failing_method} failed")

    def fail(records, output_path):
        assert runner.last_quality_result is completed_results[0]
        assert records is completed_results[0].clean_records
        raise error

    methods = ("to_csv", "to_json", "to_excel")
    for method in methods:
        original_export = getattr(runner.exporter, method)
        setattr(runner.exporter, method, Mock(wraps=original_export))
    getattr(runner.exporter, failing_method).side_effect = fail

    with pytest.raises(OSError) as caught:
        runner.run(path, tmp_path, clear_cache=True)

    assert caught.value is error
    runner.quality_processor.process.assert_called_once()
    assert runner.last_quality_result is completed_results[0]
    assert runner.last_quality_result is not previous_result
    assert runner.last_quality_result.clean_records == [{"id": 1, "title": "One", "price": 12.5}]
    assert runner.last_quality_result.rejected_records == []
    assert runner.last_quality_result.report.exported == 1
    runner.last_quality_result.report.enforce_invariants()
    failed_index = methods.index(failing_method)
    for index, method in enumerate(methods):
        if index <= failed_index:
            getattr(runner.exporter, method).assert_called_once()
        else:
            getattr(runner.exporter, method).assert_not_called()



def test_quality_audits_use_canonical_paths_same_result_and_precede_all_exports(
    tmp_path, make_quality_runner, monkeypatch,
):
    from dataclasses import asdict
    import core.exporter as exporter_module

    first = "https://example.com/page-1.html"
    second = "https://example.com/page-2.html"
    runner, path = make_quality_runner(html_by_url={first: "first", second: "second"})
    profile = runner.loader.load(path)
    profile["site_name"] = "Client / Menu: 2026"
    profile["fields"]["title"]["required"] = True
    runner.loader = FakeLoader(profile)
    records = {
        first: [
            {"id": "1", "title": "Caf\u00e9", "price": "12.50", "_source_url": first, "_page": 1},
            {"id": "bad", "title": None, "price": "1.00", "_source_url": first, "_page": 1},
        ],
        second: [{"id": "1", "title": "Duplicate", "price": "15.00"}],
    }
    original_profile, original_records = deepcopy(profile), deepcopy(records)
    runner.parser = FakeParser(records)
    runner.paginator = FakePaginator({first: second, second: None})
    runner.quality_processor.process = Mock(wraps=runner.quality_processor.process)
    snapshots = []
    original_write = runner.exporter.write_quality_artifacts

    def write(result, output_path):
        assert result is runner.last_quality_result
        snapshots.append(deepcopy(result))
        return original_write(result, output_path)

    runner.exporter.write_quality_artifacts = Mock(side_effect=write)
    dump = Mock(wraps=exporter_module.json.dump)
    monkeypatch.setattr(exporter_module.json, "dump", dump)
    output_dir = tmp_path / "client delivery"
    quality_path = output_dir / "client_menu_2026.quality.json"
    rejected_path = output_dir / "client_menu_2026.rejected.json"

    def observe_export(method):
        def export(clean_records, output_path):
            result = runner.last_quality_result
            assert clean_records is result.clean_records
            assert json.loads(quality_path.read_text(encoding="utf-8")) == asdict(result.report)
            assert json.loads(rejected_path.read_text(encoding="utf-8")) == result.rejected_records
            return method(clean_records, output_path)
        return export

    for name in ("to_csv", "to_json", "to_excel"):
        setattr(runner.exporter, name, observe_export(getattr(runner.exporter, name)))

    summary = runner.run(path, output_dir)

    result = runner.last_quality_result
    runner.quality_processor.process.assert_called_once()
    runner.exporter.write_quality_artifacts.assert_called_once_with(result, summary["csv_path"])
    assert dump.call_count == 2
    assert dump.call_args_list[1].args[0] is result.rejected_records
    assert result == snapshots[0]
    assert profile == original_profile
    assert records == original_records
    assert result.report.extracted == 3
    assert result.report.rejected == result.report.duplicates_removed == result.report.exported == 1
    assert len(result.rejected_records[0]["reasons"]) == 2
    assert result.rejected_records[0]["_source_url"] == first
    assert result.rejected_records[0]["_page"] == 1
    assert set(output_dir.iterdir()) == {
        summary["csv_path"], summary["json_path"], summary["xlsx_path"], quality_path, rejected_path,
    }


@pytest.mark.parametrize("case", ["no_rejections", "zero", "all_rejected"])
@pytest.mark.parametrize("existing_normal_outputs", [False, True])
def test_quality_audits_exist_for_empty_and_nonempty_completed_runs(
    tmp_path, make_quality_runner, case, existing_normal_outputs,
):
    from dataclasses import asdict

    runner, path = make_quality_runner()
    records = {
        "no_rejections": [{"id": "1", "title": "One", "price": "1.0"}],
        "zero": [],
        "all_rejected": [{"id": "invalid", "title": "Rejected", "price": "1.0"}],
    }[case]
    runner.parser = FakeParser({"https://example.com/page-1.html": records})
    if existing_normal_outputs:
        for extension in ("csv", "json", "xlsx"):
            (tmp_path / f"quality_example.{extension}").write_text("previous", encoding="utf-8")

    runner.run(path, tmp_path)

    result = runner.last_quality_result
    report = json.loads((tmp_path / "quality_example.quality.json").read_text(encoding="utf-8"))
    rejected = json.loads((tmp_path / "quality_example.rejected.json").read_text(encoding="utf-8"))
    assert report == asdict(result.report)
    assert rejected == result.rejected_records
    if case == "zero":
        assert all(value == 0 for value in report.values())
        assert rejected == []
        assert len(runner.exporter.calls) == 3
        assert all(call[1] == [] for call in runner.exporter.calls)
    elif case == "all_rejected":
        assert report["rejected"] == len(rejected) == 1
        assert report["exported"] == 0
        assert len(runner.exporter.calls) == 3
        assert all(call[1] == [] for call in runner.exporter.calls)
    else:
        assert rejected == []
        assert report["exported"] == 1


@pytest.mark.parametrize("quality", [None, False])
@pytest.mark.parametrize("existing", [False, True])
def test_legacy_runs_never_touch_quality_artifacts(tmp_path, make_quality_runner, quality, existing):
    runner, path = make_quality_runner(quality=quality)
    targets = [tmp_path / f"quality_example.{kind}.json" for kind in ("quality", "rejected")]
    if existing:
        for target in targets:
            target.write_bytes(b'{"previous": true}')
    runner.exporter.write_quality_artifacts = Mock(side_effect=AssertionError("Audit must not run"))

    runner.run(path, tmp_path)

    runner.exporter.write_quality_artifacts.assert_not_called()
    assert runner.last_quality_result is None
    for target in targets:
        if existing:
            assert target.read_bytes() == b'{"previous": true}'
        else:
            assert not target.exists()


@pytest.mark.parametrize("existing", [False, True])
def test_completion_preserves_audits_or_rebuilds_missing_artifacts(
    tmp_path, make_quality_runner, existing,
):
    runner, path = make_quality_runner()
    runner.parser = FakeParser({
        "https://example.com/page-1.html": [{"id": "bad", "title": "Rejected", "price": "1.0"}],
    })
    runner.run(path, tmp_path)
    targets = [tmp_path / f"quality_example.{kind}.json" for kind in ("quality", "rejected")]
    previous = [target.read_bytes() for target in targets]
    assert runner.last_quality_result.report.rejected == 1
    if not existing:
        for target in targets:
            target.unlink()
    runner.client = FakeClient(runner.client.html_by_url)
    runner.exporter.calls.clear()
    runner.exporter.write_quality_artifacts = Mock(wraps=runner.exporter.write_quality_artifacts)

    summary = runner.run(path, tmp_path)

    assert summary["cache_only_run"] is existing
    assert [target.read_bytes() for target in targets] == previous
    if existing:
        assert runner.last_quality_result is None
        assert runner.exporter.calls == []
        runner.exporter.write_quality_artifacts.assert_not_called()
    else:
        assert runner.last_quality_result.report.rejected == 1
        assert runner.last_quality_result.report.extracted == 1
        assert len(runner.exporter.calls) == 3
        runner.exporter.write_quality_artifacts.assert_called_once()


@pytest.mark.parametrize("existing", [False, True])
def test_quality_failure_writes_no_audits_or_data(tmp_path, make_quality_runner, existing):
    runner, path = make_quality_runner()
    targets = [tmp_path / f"quality_example.{kind}.json" for kind in ("quality", "rejected")]
    if existing:
        for target in targets:
            target.write_text("[]", encoding="utf-8")
    error = RuntimeError("quality failed")
    runner.quality_processor.process = Mock(side_effect=error)
    runner.exporter.write_quality_artifacts = Mock(side_effect=AssertionError("Audit must not run"))

    with pytest.raises(RuntimeError) as caught:
        runner.run(path, tmp_path)

    assert caught.value is error
    assert runner.last_quality_result is None
    runner.exporter.write_quality_artifacts.assert_not_called()
    assert runner.exporter.calls == []
    for target in targets:
        if existing:
            assert target.read_text(encoding="utf-8") == "[]"
        else:
            assert not target.exists()


@pytest.mark.parametrize("failed_write", [1, 2])
def test_audit_write_failure_preserves_target_and_prevents_normal_export(
    tmp_path, make_quality_runner, monkeypatch, failed_write,
):
    import core.exporter as exporter_module

    runner, path = make_quality_runner()
    targets = [tmp_path / f"quality_example.{kind}.json" for kind in ("quality", "rejected")]
    for target in targets:
        target.write_text('["previous"]', encoding="utf-8")
    previous = [target.read_bytes() for target in targets]
    original_dump = exporter_module.json.dump
    calls = []
    error = OSError("partial audit write failed")

    def dump(data, handle, **kwargs):
        calls.append(handle.name)
        if len(calls) == failed_write:
            handle.write('{"partial":')
            raise error
        return original_dump(data, handle, **kwargs)

    monkeypatch.setattr(exporter_module.json, "dump", dump)

    with pytest.raises(OSError) as caught:
        runner.run(path, tmp_path)

    assert caught.value is error
    assert runner.last_quality_result is not None
    assert len(calls) == failed_write
    assert runner.exporter.calls == []
    assert targets[failed_write - 1].read_bytes() == previous[failed_write - 1]
    if failed_write == 1:
        assert targets[1].read_bytes() == previous[1]
    assert not list(tmp_path.glob("*.tmp"))
    assert all(not (tmp_path / f"quality_example.{extension}").exists() for extension in ("csv", "json", "xlsx"))


@pytest.mark.parametrize("method", ["to_csv", "to_json", "to_excel"])
def test_normal_export_failure_keeps_successful_audits(tmp_path, make_quality_runner, method):
    from dataclasses import asdict

    runner, path = make_quality_runner()
    error = OSError("normal export failed")
    setattr(runner.exporter, method, Mock(side_effect=error))
    runner.exporter.write_quality_artifacts = Mock(wraps=runner.exporter.write_quality_artifacts)

    with pytest.raises(OSError) as caught:
        runner.run(path, tmp_path)

    assert caught.value is error
    result = runner.last_quality_result
    assert result is not None
    runner.exporter.write_quality_artifacts.assert_called_once()
    assert runner.exporter.write_quality_artifacts.call_args.args[0] is result
    assert json.loads((tmp_path / "quality_example.quality.json").read_text(encoding="utf-8")) == asdict(result.report)
    assert json.loads((tmp_path / "quality_example.rejected.json").read_text(encoding="utf-8")) == result.rejected_records


@pytest.fixture
def make_completion_runner(tmp_path, monkeypatch):
    from core.cache import URLCache

    cache = URLCache(str(tmp_path / "completion.db"))
    monkeypatch.setattr(runner_module, "URLCache", lambda: cache)
    monkeypatch.setattr(runner_module, "RobotsChecker", AllowAllRobots)
    urls = [f"https://example.com/page-{number}.html" for number in range(1, 4)]
    pages = {
        url: f"<article><h2>Item {number}</h2></article>" + (
            f'<a class="next" href="{urls[number]}">Next</a>' if number < 3 else ""
        )
        for number, url in enumerate(urls, 1)
    }

    def make(quality=False, output_name="outputs"):
        profile = {
            "site_name": "Completion Example", "engine": "static", "start_url": urls[0],
            "fields": {"title": "article h2::text"}, "max_pages": 3,
            "pagination": {"next_button": "a.next"}, "data_quality": quality,
        }
        path = tmp_path / "completion.yaml"
        path.write_text(yaml.safe_dump(profile, sort_keys=False), encoding="utf-8")
        runner = ScrapeRunner(client=FakeClient(pages))
        return runner, path, tmp_path / output_name, urls

    return make


def test_completed_rerun_does_not_refetch_pagination_regression(make_completion_runner):
    runner, path, output_dir, urls = make_completion_runner()
    first = runner.run(path, output_dir)
    assert runner.client.calls == urls
    before = {key: first[key].read_bytes() for key in ("csv_path", "json_path", "xlsx_path")}
    runner.client = FakeClient(runner.client.html_by_url)

    second = runner.run(path, output_dir)

    assert runner.client.calls == []
    assert second["cache_only_run"] is True
    assert {key: second[key].read_bytes() for key in before} == before


def test_mixed_legacy_cache_cannot_overwrite_complete_exports_with_subset(make_completion_runner):
    import sqlite3

    runner, path, output_dir, urls = make_completion_runner()
    first = runner.run(path, output_dir)
    expected = [{"title": f"Item {number}"} for number in range(1, 4)]
    assert json.loads(first["json_path"].read_text(encoding="utf-8")) == expected
    # Reproduce a partial legacy cache while one required output needs recovery.
    with sqlite3.connect(runner.cache.db_path) as connection:
        connection.execute("DELETE FROM url_cache")
    runner.cache.mark_done(urls[0], 1)
    first["xlsx_path"].unlink()
    runner.client = FakeClient(runner.client.html_by_url)

    rebuilt = runner.run(path, output_dir)

    assert json.loads(rebuilt["json_path"].read_text(encoding="utf-8")) == expected
    assert runner.client.calls == urls
    assert rebuilt["records_extracted"] == 3
    assert rebuilt["cached_skips"] == 0



def completion_rows(cache):
    from contextlib import closing
    import sqlite3

    with closing(sqlite3.connect(cache.db_path)) as connection:
        return connection.execute("SELECT * FROM run_completions ORDER BY output_key").fetchall()


def expected_run_paths(summary, quality):
    paths = [summary[key] for key in ("csv_path", "json_path", "xlsx_path")]
    if quality:
        paths.extend(summary["csv_path"].with_suffix(f".{kind}.json") for kind in ("quality", "rejected"))
    return paths


@pytest.mark.parametrize("quality", [False, True])
@pytest.mark.parametrize("engine", ["static", "browser"])
def test_completed_rerun_is_entirely_read_only_including_database(
    make_completion_runner, monkeypatch, quality, engine,
):
    import sqlite3
    import core.cache as cache_module

    runner, path, output_dir, urls = make_completion_runner(quality=quality)
    profile = yaml.safe_load(path.read_text(encoding="utf-8"))
    profile["engine"] = engine
    path.write_text(yaml.safe_dump(profile), encoding="utf-8")
    first = runner.run(path, output_dir)
    files = expected_run_paths(first, quality)
    before = {file: (file.read_bytes(), file.stat().st_mtime_ns) for file in files}
    database = (runner.cache.db_path.read_bytes(), runner.cache.db_path.stat().st_mtime_ns)
    database_size = runner.cache.db_path.stat().st_size
    sidecars = [Path(str(runner.cache.db_path) + suffix) for suffix in ("-journal", "-wal", "-shm")]
    sidecars_before = {file: (file.read_bytes(), file.stat().st_size, file.stat().st_mtime_ns)
                       for file in sidecars if file.exists()}
    rows = completion_rows(runner.cache)
    assert len(rows) == 1 and rows[0][2] == 3
    runner.client = FakeClient({})
    for target, methods in (
        (runner.parser, ("extract",)), (runner.transformer, ("transform",)),
        (runner.quality_processor, ("process",)),
        (runner.exporter, ("to_csv", "to_json", "to_excel", "write_quality_artifacts")),
        (runner.cache, ("invalidate_run", "mark_complete", "clear_runs", "mark_done", "clear")),
    ):
        for method in methods:
            monkeypatch.setattr(target, method, Mock(side_effect=AssertionError(f"Unexpected {method}")))
    robots = Mock(side_effect=AssertionError("No robots network request"))
    factory = Mock(side_effect=AssertionError("No browser/client construction"))
    monkeypatch.setattr(runner_module, "RobotsChecker", robots)
    monkeypatch.setattr(runner_module, "create_client", factory)
    original_connect = sqlite3.connect
    statements, total_changes = [], []

    class ObservedConnection(sqlite3.Connection):
        def close(self):
            total_changes.append(self.total_changes)
            return super().close()

    def connect(*args, **kwargs):
        connection = original_connect(*args, **kwargs, factory=ObservedConnection)
        connection.set_trace_callback(statements.append)
        return connection

    with monkeypatch.context() as patch:
        patch.setattr(cache_module.sqlite3, "connect", connect)
        summary = runner.run(path, output_dir)
    assert statements and all(statement.lstrip().upper().startswith("SELECT") for statement in statements)
    assert runner.client.calls == []
    assert runner.client.enter_count == runner.client.exit_count == 0
    assert runner.last_quality_result is None
    assert summary == {
        "site_name": "Completion Example", "pages_scraped": 0, "records_extracted": 0,
        "records_transformed": 0, "cached_skips": 3, "robots_blocked": 0,
        "cache_only_run": True, "csv_path": first["csv_path"],
        "json_path": first["json_path"], "xlsx_path": first["xlsx_path"],
    }
    assert {file: (file.read_bytes(), file.stat().st_mtime_ns) for file in files} == before
    assert (runner.cache.db_path.read_bytes(), runner.cache.db_path.stat().st_mtime_ns) == database
    assert completion_rows(runner.cache) == rows
    assert completion_rows(runner.cache)[0][3] == rows[0][3]  # completed_at
    assert runner.cache.db_path.stat().st_size == database_size
    assert total_changes == [0]
    assert {file: (file.read_bytes(), file.stat().st_size, file.stat().st_mtime_ns)
            for file in sidecars if file.exists()} == sidecars_before
    robots.assert_not_called()
    factory.assert_not_called()


def test_completion_is_last_and_identity_uses_one_expanded_profile_snapshot(
    make_completion_runner, monkeypatch,
):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    loaded = runner.loader.load(path)
    before = deepcopy(loaded)
    runner.loader.load = Mock(return_value=loaded)
    identity = Mock(wraps=runner_module.run_identity)
    monkeypatch.setattr(runner_module, "run_identity", identity)
    events = []
    for target, methods in (
        (runner.client, ("fetch",)), (runner.parser, ("extract",)),
        (runner.transformer, ("transform",)), (runner.quality_processor, ("process",)),
        (runner.exporter, ("write_quality_artifacts", "to_csv", "to_json", "to_excel")),
        (runner.cache, ("invalidate_run", "mark_complete")),
    ):
        for name in methods:
            original = getattr(target, name)

            def call(*args, _name=name, _original=original, **kwargs):
                events.append(_name)
                if _name != "invalidate_run":
                    assert completion_rows(runner.cache) == []
                if _name == "mark_complete":
                    assert all(file.stat().st_size > 0 for file in output_dir.iterdir())
                    assert len(list(output_dir.iterdir())) == 5
                return _original(*args, **kwargs)

            setattr(target, name, Mock(side_effect=call))
    summary = runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert summary["records_extracted"] == 3
    assert events == ["invalidate_run", "fetch", "extract", "fetch", "extract", "fetch", "extract",
                      "transform", "process", "write_quality_artifacts", "to_csv", "to_json",
                      "to_excel", "mark_complete"]
    runner.loader.load.assert_called_once_with(path)
    identity.assert_called_once_with(loaded, summary["csv_path"])
    assert identity.call_args.args[0] is loaded
    assert loaded == before
    assert loaded["fields"]["title"] == {
        "selector": "article h2::text", "required": False, "type": "string",
    }
    assert completion_rows(runner.cache)[0][2] == 3


@pytest.mark.parametrize("suffix", [".csv", ".json", ".xlsx", ".quality.json", ".rejected.json"])
@pytest.mark.parametrize("damage", ["missing", "empty"])
def test_incomplete_output_rebuilds_every_page_and_all_required_files(
    make_completion_runner, suffix, damage,
):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    damaged = first["csv_path"].with_suffix(suffix)
    if damage == "missing":
        damaged.unlink()
    else:
        damaged.write_bytes(b"")
    runner.client = FakeClient(runner.client.html_by_url)
    original_fetch = runner.client.fetch

    def fetch(url, wait_for=None):
        assert completion_rows(runner.cache) == []  # visible from a separate connection
        return original_fetch(url, wait_for=wait_for)

    runner.client.fetch = fetch
    rebuilt = runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert rebuilt["cache_only_run"] is False
    assert rebuilt["records_extracted"] == rebuilt["records_transformed"] == 3
    assert rebuilt["cached_skips"] == 0
    assert all(file.is_file() and file.stat().st_size > 0 for file in expected_run_paths(rebuilt, True))
    assert len(json.loads(rebuilt["json_path"].read_text(encoding="utf-8"))) == 3
    assert runner.last_quality_result.report.extracted == 3
    assert completion_rows(runner.cache)[0][2] == 3


@pytest.mark.parametrize("quality", [False, True])
def test_completion_without_any_outputs_rebuilds(make_completion_runner, quality):
    runner, path, output_dir, urls = make_completion_runner(quality=quality)
    first = runner.run(path, output_dir)
    for file in expected_run_paths(first, quality):
        file.unlink()
    runner.client = FakeClient(runner.client.html_by_url)
    second = runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert second["records_extracted"] == 3
    assert all(file.is_file() and file.stat().st_size > 0 for file in expected_run_paths(second, quality))


@pytest.mark.parametrize("suffix", [".csv", ".json", ".xlsx", ".quality.json", ".rejected.json"])
def test_directory_at_output_path_is_not_completion_evidence(make_completion_runner, suffix):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    target = first["csv_path"].with_suffix(suffix)
    target.unlink()
    target.mkdir()
    runner.client = FakeClient(runner.client.html_by_url)
    with pytest.raises(OSError):
        runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert completion_rows(runner.cache) == []
    assert target.is_dir()  # Never delete a user's directory to make room for a file.
    target.rmdir()
    runner.client = FakeClient(runner.client.html_by_url)
    runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert target.is_file()


def test_legacy_rows_are_preserved_ignored_and_never_written_by_runner(make_completion_runner):
    import sqlite3

    runner, path, output_dir, urls = make_completion_runner()
    for url in urls:
        runner.cache.mark_done(url, 99)
    with sqlite3.connect(runner.cache.db_path) as connection:
        previous = connection.execute("SELECT * FROM url_cache ORDER BY url").fetchall()
    runner.cache.mark_done = Mock(side_effect=AssertionError("No per-page writes"))
    runner.cache.is_cached = Mock(side_effect=AssertionError("No per-page skips"))
    runner.run(path, output_dir)
    assert runner.client.calls == urls
    runner.client = FakeClient({})
    assert runner.run(path, output_dir)["cache_only_run"] is True
    assert runner.client.calls == []
    with sqlite3.connect(runner.cache.db_path) as connection:
        assert connection.execute("SELECT * FROM url_cache ORDER BY url").fetchall() == previous
    runner.cache.mark_done.assert_not_called()
    runner.cache.is_cached.assert_not_called()


@pytest.mark.parametrize("stage", ["fetch", "extract", "transform", "process", "write_quality_artifacts",
                                    "to_csv", "to_json", "to_excel"])
def test_failed_stage_never_completes_and_next_invocation_rebuilds(make_completion_runner, stage):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    target = {
        "fetch": runner.client, "extract": runner.parser, "transform": runner.transformer,
        "process": runner.quality_processor,
    }.get(stage, runner.exporter)
    original = getattr(target, stage)
    error = RuntimeError(f"{stage} failed")
    setattr(target, stage, Mock(side_effect=error))
    normal_methods = ("to_csv", "to_json", "to_excel")
    for method in normal_methods:
        if method != stage:
            setattr(runner.exporter, method, Mock(wraps=getattr(runner.exporter, method)))
    with pytest.raises(RuntimeError) as caught:
        runner.run(path, output_dir)
    assert caught.value is error
    assert completion_rows(runner.cache) == []
    if stage not in normal_methods:
        for method in normal_methods:
            getattr(runner.exporter, method).assert_not_called()
    setattr(target, stage, original)
    runner.client = FakeClient(runner.client.html_by_url)
    summary = runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert summary["records_extracted"] == 3
    assert len(completion_rows(runner.cache)) == 1


@pytest.mark.parametrize("error_type", [RuntimeError, KeyboardInterrupt])
def test_interrupted_recovery_commits_invalidation_and_next_run_fully_rebuilds(
    make_completion_runner, error_type,
):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    first["xlsx_path"].write_bytes(b"")
    runner.client = FakeClient(runner.client.html_by_url)
    original_extract = runner.parser.extract
    error = error_type("interrupted after first fetch")

    def abort(*args):
        assert runner.client.calls == [urls[0]]
        assert completion_rows(runner.cache) == []
        # Even if the missing file reappears, the deleted marker must not revive.
        first["xlsx_path"].write_bytes(b"non-empty old output")
        raise error

    runner.parser.extract = abort
    with pytest.raises(error_type) as caught:
        runner.run(path, output_dir)
    assert caught.value is error
    assert completion_rows(runner.cache) == []
    assert runner.client.closed
    runner.parser.extract = original_extract
    runner.client = FakeClient(runner.client.html_by_url)
    runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert runner.last_quality_result.report.extracted == 3


def test_invalidation_failure_prevents_all_network_and_output_work(make_completion_runner, monkeypatch):
    import sqlite3

    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    first["json_path"].unlink()
    before = {file: file.read_bytes() for file in output_dir.iterdir()}
    rows = completion_rows(runner.cache)
    with sqlite3.connect(runner.cache.db_path) as connection:
        connection.execute("""CREATE TRIGGER deny_invalidation BEFORE DELETE ON run_completions
                              BEGIN SELECT RAISE(ABORT, 'invalidation failed'); END""")
    runner.client = FakeClient({})
    robots = Mock(side_effect=AssertionError("Network must not begin"))
    monkeypatch.setattr(runner_module, "RobotsChecker", robots)
    for method in ("to_csv", "to_json", "to_excel", "write_quality_artifacts"):
        setattr(runner.exporter, method, Mock(side_effect=AssertionError("No output writes")))
    with pytest.raises(sqlite3.IntegrityError, match="invalidation failed"):
        runner.run(path, output_dir)
    assert runner.client.calls == []
    assert runner.client.enter_count == 0
    assert runner.last_quality_result is None
    robots.assert_not_called()
    assert completion_rows(runner.cache) == rows
    assert {file: file.read_bytes() for file in output_dir.iterdir()} == before


def test_completion_write_failure_keeps_outputs_but_next_run_rebuilds(make_completion_runner):
    import sqlite3

    runner, path, output_dir, urls = make_completion_runner(quality=True)
    with sqlite3.connect(runner.cache.db_path) as connection:
        connection.execute("""CREATE TRIGGER deny_completion BEFORE INSERT ON run_completions
                              BEGIN SELECT RAISE(ABORT, 'completion failed'); END""")
    with pytest.raises(sqlite3.IntegrityError, match="completion failed"):
        runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert completion_rows(runner.cache) == []
    assert runner.last_quality_result.report.extracted == 3
    assert len(list(output_dir.iterdir())) == 5
    assert all(file.is_file() and file.stat().st_size > 0 for file in output_dir.iterdir())
    with sqlite3.connect(runner.cache.db_path) as connection:
        connection.execute("DROP TRIGGER deny_completion")
    runner.client = FakeClient(runner.client.html_by_url)
    runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert len(completion_rows(runner.cache)) == 1


def test_distinct_destinations_do_not_share_completion(make_completion_runner):
    runner, path, output_dir, urls = make_completion_runner()
    runner.run(path, output_dir)
    runner.client = FakeClient(runner.client.html_by_url)
    other = output_dir.parent / "different"
    runner.run(path, other)
    assert runner.client.calls == urls
    assert len(completion_rows(runner.cache)) == 2
    runner.client = FakeClient({})
    assert runner.run(path, output_dir)["cache_only_run"] is True
    assert runner.run(path, other)["cache_only_run"] is True
    assert runner.client.calls == []


@pytest.mark.parametrize("second_run_fails", [False, True])
def test_profile_switch_cannot_reuse_marker_for_overwritten_outputs(make_completion_runner, second_run_fails):
    runner, path, output_dir, urls = make_completion_runner()
    runner.run(path, output_dir)
    original_profile = path.read_text(encoding="utf-8")
    profile = yaml.safe_load(original_profile)
    profile["max_pages"] = 1
    path.write_text(yaml.safe_dump(profile), encoding="utf-8")
    runner.client = FakeClient(runner.client.html_by_url)
    original_excel = runner.exporter.to_excel
    if second_run_fails:
        runner.exporter.to_excel = Mock(side_effect=RuntimeError("export failed"))
        with pytest.raises(RuntimeError, match="export failed"):
            runner.run(path, output_dir)
        assert completion_rows(runner.cache) == []
    else:
        runner.run(path, output_dir)
    runner.exporter.to_excel = original_excel
    path.write_text(original_profile, encoding="utf-8")
    runner.client = FakeClient(runner.client.html_by_url)
    result = runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert result["records_extracted"] == 3


def test_reordered_loaded_profile_skips_but_changed_selector_rebuilds(make_completion_runner):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    runner.run(path, output_dir)
    profile = runner.loader.load(path)
    runner.loader = FakeLoader(dict(reversed(list(profile.items()))))
    runner.client = FakeClient({})
    assert runner.run(path, output_dir)["cache_only_run"] is True
    changed = deepcopy(profile)
    changed["fields"]["title"]["selector"] = "article > h2::text"
    runner.loader = FakeLoader(changed)
    runner.client = FakeClient({url: "<article><h2>Updated</h2></article>" for url in urls})
    assert runner.run(path, output_dir)["cache_only_run"] is False
    assert runner.client.calls == [urls[0]]


def test_robots_blocked_run_does_not_record_completion(make_completion_runner, monkeypatch):
    runner, path, output_dir, urls = make_completion_runner()
    blocked = Mock()
    blocked.is_allowed.return_value = False
    monkeypatch.setattr(runner_module, "RobotsChecker", Mock(return_value=blocked))
    result = runner.run(path, output_dir)
    assert result["robots_blocked"] == 1
    assert completion_rows(runner.cache) == []
    monkeypatch.setattr(runner_module, "RobotsChecker", AllowAllRobots)
    runner.client = FakeClient(runner.client.html_by_url)
    runner.run(path, output_dir)
    assert runner.client.calls == urls


def test_three_page_fresh_skip_recovery_demonstration(make_completion_runner):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    fetches = [len(runner.client.calls)]
    runner.client = FakeClient(runner.client.html_by_url)
    second = runner.run(path, output_dir)
    fetches.append(len(runner.client.calls))
    first["csv_path"].unlink()
    runner.client = FakeClient(runner.client.html_by_url)
    third = runner.run(path, output_dir)
    fetches.append(len(runner.client.calls))
    assert fetches == [3, 0, 3]
    assert second["cache_only_run"] is True
    assert third["records_extracted"] == third["records_transformed"] == 3
    assert len(json.loads(third["json_path"].read_text(encoding="utf-8"))) == 3
    assert all(file.is_file() and file.stat().st_size > 0 for file in expected_run_paths(third, True))
    runner.client = FakeClient(runner.client.html_by_url)
    fourth = runner.run(path, output_dir)
    fetches.append(len(runner.client.calls))
    assert fourth["cache_only_run"] is True
    assert fetches == [3, 0, 3, 0]
    print("Three-page demonstration: fetches = 3, 0, 3, 0; all five outputs restored; 3 records.")


@pytest.mark.parametrize("quality", [False, True])
def test_genuine_empty_run_replaces_old_outputs_then_completed_rerun_skips(make_completion_runner, quality):
    runner, path, output_dir, urls = make_completion_runner(quality=quality)
    first = runner.run(path, output_dir)
    runner.client = FakeClient({urls[0]: "<html>No records</html>"})
    second = runner.run(path, output_dir, clear_cache=True)
    assert second["records_extracted"] == second["records_transformed"] == 0
    assert second["pages_scraped"] == 1
    assert json.loads(second["json_path"].read_text(encoding="utf-8")) == []
    assert second["csv_path"].read_text(encoding="utf-8").strip() == ""
    workbook = load_workbook(second["xlsx_path"])
    try:
        assert list(workbook.active.values) == []
    finally:
        workbook.close()
    if quality:
        report = json.loads(second["csv_path"].with_suffix(".quality.json").read_text(encoding="utf-8"))
        assert all(value == 0 for value in report.values())
        assert json.loads(second["csv_path"].with_suffix(".rejected.json").read_text(encoding="utf-8")) == []
    runner.client = FakeClient({})
    third = runner.run(path, output_dir)
    assert third["cache_only_run"] is True
    assert third["cached_skips"] == 1
    assert runner.client.calls == []
    assert runner.last_quality_result is None


def test_nonempty_corrupt_outputs_are_not_content_validated(make_completion_runner):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    for file in expected_run_paths(first, True):
        file.write_bytes(b"nonempty content is intentionally not validated")
    runner.client = FakeClient({})
    assert runner.run(path, output_dir)["cache_only_run"] is True
    assert runner.client.calls == []
    assert all(file.read_bytes() == b"nonempty content is intentionally not validated"
               for file in expected_run_paths(first, True))


def test_profile_file_changes_during_fetch_do_not_change_current_snapshot(make_completion_runner):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    original_profile = runner.loader.load(path)
    original_fetch = runner.client.fetch
    original_process = runner.quality_processor.process
    runner.loader.load = Mock(wraps=runner.loader.load)
    observed = []

    def fetch(url, wait_for=None):
        path.write_text("invalid: new profile for a later invocation", encoding="utf-8")
        return original_fetch(url, wait_for=wait_for)

    def process(records, profile):
        assert profile == original_profile
        observed.append(profile)
        return original_process(records, profile)

    runner.client.fetch = fetch
    runner.quality_processor.process = process
    result = runner.run(path, output_dir)
    runner.loader.load.assert_called_once_with(path)
    assert len(observed) == 1
    assert runner.client.calls == urls
    output_key, fingerprint = runner_module.run_identity(original_profile, result["csv_path"])
    assert runner.cache.completed_pages(output_key, fingerprint) == 3


def test_exporter_that_leaves_empty_required_output_cannot_mark_completion(make_completion_runner):
    runner, path, output_dir, urls = make_completion_runner()
    original_json = runner.exporter.to_json

    def empty_json(records, output_path):
        output_path.write_bytes(b"")
        return output_path

    runner.exporter.to_json = empty_json
    with pytest.raises(RuntimeError, match="missing or empty"):
        runner.run(path, output_dir)
    assert completion_rows(runner.cache) == []
    runner.exporter.to_json = original_json
    runner.client = FakeClient(runner.client.html_by_url)
    runner.run(path, output_dir)
    assert runner.client.calls == urls



def test_explicit_clear_cache_preserves_historical_full_clear_scope(make_completion_runner, monkeypatch):
    import sqlite3

    runner, path, output_dir, urls = make_completion_runner(quality=True)
    runner.run(path, output_dir)
    runner.cache.mark_done(urls[0], 99)
    runner.cache.mark_complete("unrelated-output", "other-profile", 2)
    pages = runner.client.html_by_url
    fresh_client = FakeClient(pages)
    runner.client = None

    def create(engine):
        assert completion_rows(runner.cache) == []
        with sqlite3.connect(runner.cache.db_path) as connection:
            assert connection.execute("SELECT COUNT(*) FROM url_cache").fetchone() == (0,)
        return fresh_client

    factory = Mock(side_effect=create)
    monkeypatch.setattr(runner_module, "create_client", factory)
    summary = runner.run(path, output_dir, clear_cache=True)

    factory.assert_called_once_with("static")
    assert fresh_client.calls == urls
    assert summary["cache_only_run"] is False
    assert len(completion_rows(runner.cache)) == 1
    assert not runner.cache.is_cached(urls[0])


def test_recovery_invalidation_is_committed_before_client_and_preserves_other_jobs(
    make_completion_runner, monkeypatch,
):
    runner, path, output_dir, urls = make_completion_runner()
    first = runner.run(path, output_dir)
    runner.cache.mark_done(urls[0], 99)
    runner.cache.mark_complete("unrelated-output", "other-profile", 2)
    other = [row for row in completion_rows(runner.cache) if row[0] == "unrelated-output"]
    first["json_path"].unlink()
    fresh_client = FakeClient(runner.client.html_by_url)
    runner.client = None
    runner.cache.clear = Mock(side_effect=AssertionError("Recovery must not clear all caches"))

    def create(engine):
        assert completion_rows(runner.cache) == other
        assert runner.cache.is_cached(urls[0])
        return fresh_client

    factory = Mock(side_effect=create)
    monkeypatch.setattr(runner_module, "create_client", factory)
    runner.run(path, output_dir)

    factory.assert_called_once_with("static")
    assert fresh_client.calls == urls
    assert [row for row in completion_rows(runner.cache) if row[0] == "unrelated-output"] == other
    assert runner.cache.is_cached(urls[0])
    runner.cache.clear.assert_not_called()


@pytest.mark.parametrize("setting, value", [
    (("site_name",), "Different Output"), (("engine",), "browser"),
    (("start_url",), "https://example.com/new-start.html"), (("max_pages",), 2),
    (("pagination", "next_button"), "a.other"), (("delay",), 1), (("wait_for",), "article"),
    (("fields", "title", "selector"), "article h2:not(.excluded)::text"),
    (("fields", "title", "required"), True), (("fields", "title", "type"), "URL"),
    (("fields", "date", "format"), "%d/%m/%Y"), (("data_quality",), False),
    (("unique_key",), ["title"]), (("duplicate_policy",), "report_only"),
    (("include_provenance",), True), (("record_selector",), "article"),
    (("transform",), {"title": "strip"}), (("transformation",), {"title": "strip"}),
    (("transformations",), {"title": "strip"}),
])
def test_each_included_profile_category_forces_actual_rebuild(
    make_completion_runner, monkeypatch, setting, value,
):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    profile = yaml.safe_load(path.read_text(encoding="utf-8"))
    profile["fields"]["title"] = {"selector": "article h2::text"}
    profile["fields"]["date"] = {"selector": "article time::text", "type": "date", "format": "%Y-%m-%d"}
    path.write_text(yaml.safe_dump(profile), encoding="utf-8")
    runner.run(path, output_dir)
    target = profile
    for key in setting[:-1]:
        target = target[key]
    target[setting[-1]] = value
    if setting == ("data_quality",):
        profile["fields"] = {name: definition["selector"] for name, definition in profile["fields"].items()}
    path.write_text(yaml.safe_dump(profile), encoding="utf-8")
    pages = dict(runner.client.html_by_url)
    pages[profile["start_url"]] = pages[urls[0]]
    runner.client = FakeClient(pages)
    monkeypatch.setattr(runner_module.time, "sleep", Mock())

    result = runner.run(path, output_dir)

    assert result["cache_only_run"] is False
    assert runner.client.calls and runner.client.calls[0] == profile["start_url"]
    assert result["pages_scraped"] == len(runner.client.calls)


@pytest.mark.parametrize("key, value", [
    ("description", "Reworded description"), ("owner", "Delivery team"),
    ("tags", ["portfolio", "reviewed"]), ("notes", {"revision": 2}),
])
def test_excluded_behavior_neutral_metadata_does_not_rebuild(make_completion_runner, key, value):
    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    before = {file: file.read_bytes() for file in expected_run_paths(first, True)}
    rows = completion_rows(runner.cache)
    profile = yaml.safe_load(path.read_text(encoding="utf-8"))
    profile[key] = value
    path.write_text(yaml.safe_dump(profile), encoding="utf-8")
    runner.client = FakeClient({})

    result = runner.run(path, output_dir)

    assert result["cache_only_run"] is True
    assert runner.client.calls == []
    assert completion_rows(runner.cache) == rows
    assert {file: file.read_bytes() for file in before} == before


@pytest.mark.parametrize("failed_artifact", [1, 2])
def test_either_audit_file_failure_invalidates_completion_and_next_run_rebuilds(
    make_completion_runner, monkeypatch, failed_artifact,
):
    import core.exporter as exporter_module

    runner, path, output_dir, urls = make_completion_runner(quality=True)
    first = runner.run(path, output_dir)
    first["csv_path"].unlink()
    runner.client = FakeClient(runner.client.html_by_url)
    original_dump = exporter_module.json.dump
    calls = []
    error = OSError("audit artifact write failed")

    def fail(data, handle, **kwargs):
        calls.append(handle.name)
        if len(calls) == failed_artifact:
            handle.write("partial")
            raise error
        return original_dump(data, handle, **kwargs)

    with monkeypatch.context() as patch:
        patch.setattr(exporter_module.json, "dump", fail)
        with pytest.raises(OSError) as caught:
            runner.run(path, output_dir)
    assert caught.value is error
    assert len(calls) == failed_artifact
    assert completion_rows(runner.cache) == []
    assert not first["csv_path"].exists()
    runner.client = FakeClient(runner.client.html_by_url)
    runner.run(path, output_dir)
    assert runner.client.calls == urls
    assert runner.last_quality_result.report.extracted == 3
    assert len(completion_rows(runner.cache)) == 1


def test_legacy_only_database_migrates_then_rebuilds_skips_and_full_clears(tmp_path, monkeypatch):
    from contextlib import closing
    import sqlite3
    from core.cache import URLCache

    db = tmp_path / "old.db"
    url = "https://example.com/old"
    with closing(sqlite3.connect(db)) as connection, connection:
        connection.execute("""CREATE TABLE url_cache (
            url TEXT PRIMARY KEY, record_count INTEGER NOT NULL,
            scraped_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP)""")
        connection.execute("INSERT INTO url_cache VALUES (?, 99, '2000-01-01')", (url,))
        connection.execute("CREATE TABLE unrelated (value TEXT)")
        connection.execute("INSERT INTO unrelated VALUES ('retain')")
    cache = URLCache(str(db))
    monkeypatch.setattr(runner_module, "URLCache", lambda: cache)
    monkeypatch.setattr(runner_module, "RobotsChecker", AllowAllRobots)
    path = tmp_path / "legacy.yaml"
    profile = {"site_name": "Migration", "engine": "static", "start_url": url,
               "fields": {"title": "article h2::text"}}
    path.write_text(yaml.safe_dump(profile), encoding="utf-8")
    runner = ScrapeRunner(client=FakeClient({url: "<article><h2>Current</h2></article>"}))
    assert cache.is_cached(url)
    assert completion_rows(cache) == []
    first = runner.run(path, tmp_path / "outputs")
    assert runner.client.calls == [url]
    assert first["records_extracted"] == 1
    assert cache.is_cached(url)
    assert len(completion_rows(cache)) == 1
    runner.client = FakeClient({})
    assert runner.run(path, tmp_path / "outputs")["cache_only_run"] is True
    assert runner.client.calls == []
    cache.clear()
    assert not cache.is_cached(url)
    assert completion_rows(cache) == []
    with closing(sqlite3.connect(db)) as connection:
        assert connection.execute("SELECT * FROM unrelated").fetchall() == [("retain",)]


def test_zero_record_files_are_nonempty_and_identical_rerun_is_read_only(
    make_completion_runner, monkeypatch,
):
    import sqlite3
    import core.cache as cache_module

    runner, path, output_dir, urls = make_completion_runner(quality=True)
    runner.client = FakeClient({urls[0]: "<html>No listings</html>"})
    first = runner.run(path, output_dir)
    files = expected_run_paths(first, True)
    sizes = {file.name: file.stat().st_size for file in files}
    assert first["records_extracted"] == first["records_transformed"] == 0
    assert all(file.is_file() and file.stat().st_size > 0 for file in files)
    rows = completion_rows(runner.cache)
    before = {file: (file.read_bytes(), file.stat().st_size, file.stat().st_mtime_ns)
              for file in files + [runner.cache.db_path]}
    sidecars = [Path(str(runner.cache.db_path) + suffix) for suffix in ("-journal", "-wal", "-shm")]
    sidecars_before = {file: (file.read_bytes(), file.stat().st_size, file.stat().st_mtime_ns)
                       for file in sidecars if file.exists()}
    original_connect = sqlite3.connect
    statements, totals = [], []

    class ObservedConnection(sqlite3.Connection):
        def close(self):
            totals.append(self.total_changes)
            return super().close()

    def connect(*args, **kwargs):
        connection = original_connect(*args, **kwargs, factory=ObservedConnection)
        connection.set_trace_callback(statements.append)
        return connection

    runner.client = FakeClient({})
    for target, methods in (
        (runner.parser, ("extract",)), (runner.transformer, ("transform",)),
        (runner.quality_processor, ("process",)),
        (runner.exporter, ("to_csv", "to_json", "to_excel", "write_quality_artifacts")),
    ):
        for name in methods:
            setattr(target, name, Mock(side_effect=AssertionError("No rerun work or output writes")))
    with monkeypatch.context() as patch:
        patch.setattr(cache_module.sqlite3, "connect", connect)
        second = runner.run(path, output_dir)
    assert second["cache_only_run"] is True
    assert runner.client.calls == []
    assert runner.last_quality_result is None
    assert totals == [0]
    assert statements and all(statement.strip().upper().startswith("SELECT") for statement in statements)
    assert completion_rows(runner.cache) == rows
    assert {file: (file.read_bytes(), file.stat().st_size, file.stat().st_mtime_ns) for file in before} == before
    assert {file: (file.read_bytes(), file.stat().st_size, file.stat().st_mtime_ns)
            for file in sidecars if file.exists()} == sidecars_before
    print("Zero-record bytes:", sizes, "; identical rerun: 0 fetches, 0 writes, total_changes=0")
